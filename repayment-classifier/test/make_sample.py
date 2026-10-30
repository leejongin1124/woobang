"""가짜 고객 정보로 '변제금상환내역조회' 모양의 시험용 엑셀을 만든다. 금액은 2026-10-30 자료와 같다."""
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

# (회차, 지원구분, 원금, 이자, 기타채무)
ROWS = [
    (21, "일반지원(원금균등)", 35582, 0, 0),
    (20, "일반지원(원금균등)", 35582, 0, 0),
    (17, "금융취약계층/기초수급자", 72917, 0, 0),
    (16, "사전채무조정(일반)", 34998, 11494, 282),
    (16, "사전채무조정(일반)", 34791, 11426, 224),
    (16, "사전채무조정(일반)", 62101, 20395, 377),
    (15, "사전채무조정(일반)", 34904, 11588, 282),
    (15, "사전채무조정(일반)", 34697, 11520, 224),
    (15, "사전채무조정(일반)", 61934, 20562, 377),
    (8, "일반지원(원금균등)", 36035, 0, 0),
    (6, "신속채무조정/일반", 0, 8799, 0),
    (6, "신속채무조정/일반", 0, 10822, 0),
    (6, "신속채무조정/일반", 0, 3592, 0),
    (6, "신속채무조정/일반", 0, 15043, 0),
    (6, "사전(새출발기금)/개인사업자", 33122, 17816, 179),
    (6, "사전(새출발기금)/개인사업자", 34214, 18405, 185),
    (2, "일반지원(원금균등)", 22848, 0, 0),
    (1, "사전채무조정(일반)", 0, 3782, 0),
]
HDR = ["No.", "회차", "일시납\n감면율", "처리일자", "지점명", "고객명", "고객식별번호", "지원구분", "채무구분",
       "주채무자\n주민등록번호", "주채무자", "계좌번호", "원금", "이자", "기타채무", "합계", "확정\n채권금액",
       "변제금\n수취계좌", "지로코드"]

out = sys.argv[1]
wb = Workbook()
ws = wb.active
ws.title = "Sheet1"
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HDR))
ws.cell(1, 1, f"변제금상환내역조회(총 {len(ROWS)}건)").font = Font(bold=True, size=14)
ws.cell(1, 1).alignment = Alignment(horizontal="center")
ws.cell(2, 12, "거래코드:405002")
ws.cell(2, 16, "출력일시:2026.10.30 17:03:38")
ws.cell(2, 19, "사용자명:테스트")
thin = Side(style="thin")
fill = PatternFill("solid", fgColor="DDDDDD")
for c, h in enumerate(HDR, 1):
    cell = ws.cell(3, c, h)
    cell.fill = fill
    cell.font = Font(bold=True)
    cell.alignment = Alignment(wrap_text=True, horizontal="center")
    cell.border = Border(thin, thin, thin, thin)
for i, (hoe, kind, p, it, etc) in enumerate(ROWS, 1):
    r = 3 + i
    vals = [i, hoe, None, "2026-10-30", "본점", f"가짜{i:02d}", f"000000-0{i:06d}", kind, "주채무", None, None,
            f"M0000000000000{i:04d}", p, it, etc, p + it + etc, 1000000 + i, "000000000000", "GB00000"]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(r, c, v)
        cell.border = Border(thin, thin, thin, thin)
        if isinstance(v, int) and c >= 13:
            cell.number_format = "#,##0"
wb.save(out)
